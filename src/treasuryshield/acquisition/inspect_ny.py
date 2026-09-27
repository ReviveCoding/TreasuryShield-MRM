"""Inspect actual New York Fed files before fixing canonical field choices."""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
import pyreadstat


def main() -> None:
    root = Path(__file__).resolve().parents[3]
    base = root / "data" / "interim" / "newyorkfed"
    files = sorted(base.glob("*.dta"))
    if len(files) != 2:
        raise FileNotFoundError(f"Expected two extracted New York Fed .dta files; found {files}")
    summary = {}
    for path in files:
        _, meta = pyreadstat.read_dta(str(path), metadataonly=True)
        summary[path.name] = {
            "number_rows": meta.number_rows,
            "number_columns": meta.number_columns,
            "columns": [
                {"name": name, "label": meta.column_names_to_labels.get(name, "")}
                for name in meta.column_names
            ],
        }
        print(path.name, meta.number_rows, meta.number_columns, flush=True)
    book = openpyxl.load_workbook(
        root / "data" / "raw" / "newyorkfed" / "historical_call_data_dictionary.xlsx",
        read_only=True, data_only=True,
    )
    summary["dictionary"] = {
        "sheets": book.sheetnames,
        "first_rows": {
            sheet.title: [list(row) for row in sheet.iter_rows(min_row=1, max_row=8, values_only=True)]
            for sheet in book.worksheets
        },
    }
    out = root / "data" / "manifests" / "ny_schema.json"
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(json.dumps(summary, indent=2, default=str), encoding="utf-8")
    print("Wrote", out)


if __name__ == "__main__":
    main()
